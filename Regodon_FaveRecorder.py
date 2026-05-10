import openpyxl as katchot
import os

workbook = katchot.Workbook()
sheet = workbook.active

sheet["A1"] = "ID"
sheet["B1"] = "First Name"
sheet["C1"] = "Last Name"
sheet["D1"] = "Birth Year"
sheet["E1"] = "Age"
sheet["A2"] = 1
sheet["A3"] = 2
sheet["A4"] = 3

workbook.save("favorite_people.xlsx")

print("\nFavorite Person 1")
fst1name = input("Enter first name: ")
lst1name = input("Enter last name: ")
bday1 = int(input("Enter birth year: "))

print("\nFavorite Person 2")
fst2name = input("Enter first name: ")
lst2name = input("Enter last name: ")
bday2 = int(input("Enter birth year: "))

print("\nFavorite Person 3")
fst3name = input("Enter first name: ")
lst3name = input("Enter last name: ")
bday3 = int(input("Enter birth year: "))

print("\nFavorite people recorded successfully!")

age1 = 2026 - bday1
age2 = 2026 - bday2   
age3 = 2026 - bday3

wbk = katchot.load_workbook("favorite_people.xlsx")
sheet = wbk.active

print("\n+++ FAVORITE PEOPLE +++\n")

sheet["B2"] = fst1name
sheet["C2"] = lst1name
sheet["D2"] = bday1
sheet["E2"] = age1

sheet["B3"] = fst2name
sheet["C3"] = lst2name
sheet["D3"] = bday2
sheet["E3"] = age2

sheet["B4"] = fst3name
sheet["C4"] = lst3name
sheet["D4"] = bday3
sheet["E4"] = age3

wbk.save("favorite_people.xlsx")

for rows in sheet.iter_rows(values_only=True):
    print(rows)

print("\n")
os.system("pause")

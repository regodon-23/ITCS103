import openpyxl as op
import os

workbook = op.Workbook()
sheet = workbook.active

sheet["A1"] = "ID"
sheet["B2"] = "First Name"
sheet["C3"] = "Last Name"
sheet["D4"] = "Birth Year"
sheet["E5"] = "Age"
sheet["F6"] = 1
sheet["G7"] = 2
sheet["H8"] = 3

workbook.save("favorite_people.xlsx")

print("\nFavorite Person A")
fst1name = input("Enter first name: ")
lst1name = input("Enter last name: ")
bday1 = int(input("Enter birth year: "))

print("\nFavorite Person B")
fst2name = input("Enter first name: ")
lst2name = input("Enter last name: ")
bday2 = int(input("Enter birth year: "))

print("\nFavorite Person C")
fst3name = input("Enter first name: ")
lst3name = input("Enter last name: ")
bday3 = int(input("Enter birth year: "))

print("\nFavorite people recorded successfully!")

age1 = 2026 - bday1
age2 = 2026 - bday2   
age3 = 2026 - bday3

wbk = op.load_workbook("favorite_people.xlsx")
sheet = wbk.active

print("\n+++ FAVORITE PEOPLE +++\n")

sheet["A2"] = fst1name
sheet["B2"] = lst1name
sheet["C2"] = bday1
sheet["D2"] = age1

sheet["E3"] = fst2name
sheet["F3"] = lst2name
sheet["G3"] = bday2
sheet["H3"] = age2

sheet["I4"] = fst3name
sheet["J4"] = lst3name
sheet["K4"] = bday3
sheet["L4"] = age3

wbk.save("favorite_people.xlsx")

for rows in sheet.iter_rows(values_only=True):
    print(rows)

print("\n")
os.system("pause")

